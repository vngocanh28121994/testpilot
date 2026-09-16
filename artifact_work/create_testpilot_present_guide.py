from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/tuoiha17/projects/testpilot/docs/TestPilot_Tai_lieu_trinh_bay.docx"

NAVY = "17324D"
BLUE = "2D6A8A"
PALE_BLUE = "EAF3F8"
PALE_GREEN = "EAF5EF"
PALE_GOLD = "FBF4E4"
LIGHT = "F4F6F8"
BORDER = "D9D9D9"
MUTED = RGBColor(92, 108, 128)
GREEN = RGBColor(20, 120, 75)
RED = RGBColor(180, 55, 55)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Liberation Sans", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(10.6)
    normal.font.color.rgb = RGBColor(28, 36, 48)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Liberation Sans"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 21, 12, 9),
        ("Heading 2", 15, 10, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Liberation Sans"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(10.4)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.left_indent = Inches(0.24)
        style.paragraph_format.first_line_indent = Inches(-0.16)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("TestPilot  |  ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        add_page_number(p)


def add_overline(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text.upper())
    set_run_font(run, size=9, bold=True, color=RGBColor(45, 106, 138))


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(text)
    set_run_font(run, size=14, color=MUTED)


def add_para(doc, text, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
        run.italic = italic
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.24 + 0.22 * level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(lead)
            set_run_font(r1, bold=True)
            r2 = p.add_run(rest)
            set_run_font(r2)
        else:
            r = p.add_run(item)
            set_run_font(r)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.24)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{index}.  {item}")
        set_run_font(r)


def add_flow(doc, steps, fill=PALE_BLUE):
    cols = len(steps) * 2 - 1
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_width = 6.72
    arrow_w = 0.22
    step_w = (total_width - arrow_w * (len(steps) - 1)) / len(steps)
    for idx, step in enumerate(steps):
        cell = table.cell(0, idx * 2)
        cell.width = Inches(step_w)
        set_cell_fill(cell, fill)
        set_cell_border(cell)
        set_cell_margins(cell, 100, 80, 100, 80)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(step)
        set_run_font(run, size=8.8, bold=True)
        if idx < len(steps) - 1:
            arrow = table.cell(0, idx * 2 + 1)
            arrow.width = Inches(arrow_w)
            set_cell_border(arrow, "FFFFFF", "0")
            arrow.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p2 = arrow.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            rr = p2.add_run("→")
            set_run_font(rr, size=11, bold=True, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths=None, header_fill=NAVY, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        if widths:
            cell.width = Inches(widths[idx])
        set_cell_fill(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            if widths:
                cell.width = Inches(widths[c_idx])
            set_cell_fill(cell, "FFFFFF" if r_idx % 2 == 0 else LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_speaker_script(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    lead = p.add_run("Lời nói đề xuất  ")
    set_run_font(lead, bold=True, color=RGBColor(45, 106, 138))
    body = p.add_run(text)
    set_run_font(body)
    body.italic = True


def add_section_break(doc):
    doc.add_page_break()


doc = Document()
style_document(doc)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(42)
add_overline(doc, "Tài liệu dành cho người trình bày")
title = doc.add_paragraph(style="Title")
title.add_run("Hướng dẫn trình bày TestPilot")
add_subtitle(doc, "Device Farm Runner  Local Runner  App Studio Workflow")
add_para(
    doc,
    "Tài liệu này giúp người trình bày giới thiệu TestPilot theo một hành trình liền mạch: "
    "mở rộng độ phủ trên hạ tầng AWS, rút ngắn vòng lặp kiểm thử tại local, rồi tự động hoá toàn bộ "
    "quy trình từ tài liệu nghiệp vụ đến bằng chứng chạy test.",
)
doc.add_paragraph().paragraph_format.space_after = Pt(24)
add_table(
    doc,
    ["Đối tượng", "Kết quả sau buổi giới thiệu"],
    [
        ["QA  QA Automation  BA  Developer  Engineering Manager", "Hiểu khi nào dùng từng luồng và cách đọc kết quả"],
        ["Thời lượng đề xuất", "25 đến 35 phút gồm demo và hỏi đáp"],
        ["Thông điệp chính", "Một bộ kịch bản theo ý định có thể chạy trên web Android iOS và Device Farm"],
    ],
    widths=[2.05, 4.67],
)

add_section_break(doc)
doc.add_heading("Cách định vị ba luồng", level=1)
add_para(
    doc,
    "Ba phần của TestPilot giải quyết ba lớp nhu cầu khác nhau. Device Farm Runner mở rộng độ phủ thiết bị. "
    "Local Runner tối ưu tốc độ phản hồi và khả năng debug. App Studio Workflow điều phối toàn bộ vòng đời, "
    "bao gồm sinh kịch bản, review, chạy, healing và report.",
)
add_flow(doc, ["Device Farm Runner", "Kịch bản", "Local Runner", "App Studio Workflow"], PALE_GREEN)
add_table(
    doc,
    ["Luồng", "Câu hỏi nó trả lời", "Nên dùng khi"],
    [
        ["Device Farm Runner", "Ứng dụng có ổn trên nhiều mẫu máy và phiên bản hệ điều hành không", "Regression trước release hoặc cần bằng chứng trên thiết bị thật"],
        ["Kịch bản và Local Runner", "Lỗi nằm ở test ứng dụng dữ liệu hay môi trường", "Phát triển hằng ngày và cần sửa nhanh"],
        ["App Studio Workflow", "Có thể đi từ yêu cầu đến kết quả với ít thao tác thủ công hơn không", "Tạo hoặc cập nhật coverage từ tài liệu nghiệp vụ"],
    ],
    widths=[1.65, 2.55, 2.52],
)
doc.add_heading("Câu mở đầu đề xuất", level=2)
add_speaker_script(
    doc,
    "TestPilot không bắt người dùng chọn một cách chạy duy nhất. Cùng một bộ kịch bản được dùng ở ba mức: "
    "chạy rộng trên cloud, chạy nhanh tại local, hoặc để hệ thống tự điều phối từ tài liệu đến report. "
    "Tôi sẽ bắt đầu từ nơi có độ phủ thiết bị lớn nhất, sau đó đi dần về vòng lặp phát triển và cuối cùng là workflow tự động hoàn chỉnh.",
)
doc.add_heading("Nền tảng kiến trúc dùng chung", level=2)
add_bullets(doc, [
    ("Kịch bản theo ý định  ", "Gherkin mô tả thao tác như bấm nút hoặc kiểm tra nội dung, không nhúng selector vào testcase."),
    ("Element registry dùng chung  ", "Mỗi element có locator theo từng nền tảng, cho phép một scenario chạy qua Playwright hoặc Appium."),
    ("Resolver và executor dùng chung  ", "Tự chờ element, retry có kiểm soát, thu screenshot và phân loại lỗi theo cùng một chuẩn."),
    ("Report dùng chung  ", "Kết quả, ảnh, video, healing và trạng thái known issue được trình bày theo cùng cấu trúc."),
])

# Device Farm
add_section_break(doc)
add_overline(doc, "Phần một")
doc.add_heading("Device Farm Runner", level=1)
add_speaker_script(
    doc,
    "Nếu cần biết ứng dụng chạy thế nào trên nhiều thiết bị thật mà không muốn mua, cắm và bảo trì từng máy, "
    "Device Farm Runner là điểm bắt đầu. TestPilot chuẩn bị gói Appium, gửi ứng dụng và testcase lên AWS, "
    "theo dõi trạng thái rồi kéo toàn bộ bằng chứng về cùng giao diện lịch sử.",
)
doc.add_heading("Flow sử dụng", level=2)
add_flow(doc, ["Kiểm tra AWS", "Chọn project và pool", "Upload app và test", "AWS chạy từng device", "Thu report và media"])
add_numbered(doc, [
    "Mở Device Farm và kiểm tra trạng thái đăng nhập AWS. Nếu dùng SSO, làm mới phiên đăng nhập trước khi chạy.",
    "Chọn region us west 2, project, hệ điều hành Android hoặc iOS và device pool phù hợp.",
    "Chọn APK hoặc IPA, test package, giới hạn thời gian mỗi job và tuỳ chọn quay video.",
    "Bắt đầu run. Giao diện hiển thị bốn giai đoạn gồm đóng gói, upload, chờ AWS chạy và thu artifact.",
    "Mở chi tiết run để xem kết quả theo từng thiết bị, testcase pass, known issue, testcase fail, screenshot, video và log.",
])
doc.add_heading("Điểm cần chỉ trực tiếp trên giao diện", level=2)
add_bullets(doc, [
    "Trạng thái credential và thời hạn còn lại của phiên đăng nhập tạm thời.",
    "Bộ lọc thiết bị theo hãng, model, OS, form factor và availability.",
    "Device pool cố định bằng danh sách ARN để tránh thành viên thay đổi ngoài ý muốn.",
    "Giới hạn phút mỗi job để một suite treo không tiếp tục phát sinh chi phí.",
    "Lịch sử run và report chi tiết được kéo về local trước khi presigned URL hết hạn.",
])

add_section_break(doc)
doc.add_heading("Kiến trúc Device Farm Runner", level=1)
add_flow(doc, ["React UI", "Node API", "AWS SDK", "Device Farm Appium", "Artifact Collector"], PALE_GOLD)
add_para(
    doc,
    "Giao diện React gọi server Node trong cùng repository. Server dùng AWS SDK for JavaScript và credential chain mặc định, "
    "không nhận access key từ trình duyệt. Trước khi upload, hệ thống kiểm tra credential, file cấu hình và khả năng tương thích "
    "giữa app với device pool. Sau đó TestPilot tạo ba upload gồm app, Appium Node test package và testspec, chờ AWS xử lý xong "
    "mới schedule run.",
)
add_para(
    doc,
    "Trên host do AWS quản lý, testspec chọn Node 22 và Appium 2. Android dùng UiAutomator2; iOS dùng XCUITest cùng WebDriverAgent "
    "được AWS build sẵn. Mỗi thiết bị là một job riêng. Khi AWS hoàn tất, collector tải report, screenshot, video và dữ liệu flaky "
    "về run directory riêng, rồi liên kết chúng vào lịch sử TestPilot.",
)
doc.add_heading("AI hỗ trợ ở đâu", level=2)
add_bullets(doc, [
    ("Không dùng AI để điều khiển AWS  ", "Upload, schedule, polling và download artifact là luồng xác định bằng AWS SDK."),
    ("AI hỗ trợ khi chạy testcase  ", "Nếu locator đã biết và các chiến lược deterministic không tìm được element, semantic discovery hoặc vision có thể đề xuất candidate mới."),
    ("AI hỗ trợ chẩn đoán  ", "Rule engine phân loại lỗi trước; AI chỉ bổ sung ngữ cảnh khi dữ liệu đủ và không thay đổi verdict nghiệp vụ."),
    ("Healing có chốt chặn  ", "Candidate phải được verify với label và trạng thái UI; đề xuất được ghi vào telemetry và report để con người review."),
])
doc.add_heading("Điểm mạnh", level=2)
add_table(
    doc,
    ["Điểm mạnh", "Giá trị thực tế"],
    [
        ["Không quản lý phòng thiết bị", "AWS chịu trách nhiệm phần cứng, cấp phát máy và host automation"],
        ["Đa dạng model và OS", "Kiểm tra khác biệt theo hãng, kích thước màn hình và phiên bản hệ điều hành"],
        ["Thiết bị thật", "Phát hiện vấn đề quyền hệ điều hành, hiệu năng, WebView và hành vi phần cứng tốt hơn emulator"],
        ["Bằng chứng tập trung", "Report, screenshot, video và log được kéo về cùng lịch sử TestPilot"],
        ["Có thể phục hồi việc thu kết quả", "Nếu laptop mất mạng hoặc token hết hạn, có thể nối lại theo ARN bằng farm pull"],
    ],
    widths=[2.05, 4.67],
)

add_section_break(doc)
doc.add_heading("Hạn chế và cách sử dụng phù hợp", level=1)
add_table(
    doc,
    ["Hạn chế", "Tác động", "Cách giảm rủi ro"],
    [
        ["Tính phí theo phút thiết bị", "Pool càng lớn và suite càng dài thì chi phí càng tăng", "Chạy smoke nhỏ trước, đặt job timeout, chỉ dùng pool cần thiết"],
        ["Có thời gian queue upload và cài đặt", "Phản hồi chậm hơn local", "Dùng Local Runner để debug trước khi gửi farm"],
        ["Android và iOS là hai app package khác nhau", "Không thể dùng một upload cho cả hai OS", "Tạo run riêng theo APK và IPA, giữ chung feature"],
        ["Credential SSO có thể hết hạn", "AWS vẫn chạy nhưng client local có thể mất kết nối khi thu kết quả", "Kiểm tra thời hạn trước run và dùng farm pull để nối lại"],
        ["Hybrid app cần WebView inspectable", "Không có context WebView thì Appium chỉ thấy lớp native", "Bật WebContentsDebugging hoặc isInspectable trong build test"],
        ["Khó debug tương tác trực tiếp", "Không thuận tiện để sửa locator từng bước", "Dùng video log và ảnh để khoanh vùng rồi tái hiện tại local"],
    ],
    widths=[1.65, 2.37, 2.70],
    font_size=8.7,
)
doc.add_heading("Thông điệp chốt phần Device Farm", level=2)
add_speaker_script(
    doc,
    "Device Farm không thay Local Runner. Nó trả lời câu hỏi về độ phủ thiết bị và độ tin cậy trước release. "
    "Khi một case fail, chúng ta dùng bằng chứng từ farm để quyết định cần sửa sản phẩm, dữ liệu, môi trường hay automation; "
    "phần debug nhanh sẽ diễn ra ở Local Runner.",
)

# Local Runner
add_section_break(doc)
add_overline(doc, "Phần hai")
doc.add_heading("Kịch bản và Local Runner", level=1)
add_speaker_script(
    doc,
    "Sau khi đã thấy cách chạy rộng trên cloud, ta quay về vòng lặp hằng ngày. Người dùng chọn kịch bản đã review, "
    "chọn môi trường và nền tảng, chạy ngay trên trình duyệt hoặc thiết bị đang kết nối, xem log trực tiếp và mở report ngay khi hoàn tất.",
)
doc.add_heading("Flow sử dụng", level=2)
add_flow(doc, ["Chọn kịch bản", "Chọn nền tảng", "Preflight", "Chạy và theo dõi", "Mở report"])
add_numbered(doc, [
    "Vào Kịch bản để xem, sửa, duyệt hoặc gắn known issue cho từng scenario. Chỉ scenario đã duyệt mới được dùng làm bằng chứng coverage.",
    "Mở Local Runner, chọn web, Android hoặc iOS; chọn tag, environment, thiết bị và nguồn app nếu chạy native.",
    "Đọc kết quả preflight. Web cần base URL; Android cần ADB và Appium; iOS cần thiết bị hoặc simulator, trust, tunnel và WDA phù hợp.",
    "Bắt đầu run. Log hiển thị scenario đang chạy, locator được dùng, retry, popup hệ điều hành, lỗi môi trường và bước cuối cùng đã thực hiện.",
    "Mở report để xem tổng số scenario, pass, known issue, fail, flaky, bước chưa chạy, screenshot theo testcase, video và chapter thời gian.",
])
doc.add_heading("Nguồn app cho lượt chạy native", level=2)
add_table(
    doc,
    ["Lựa chọn", "Khi nên dùng", "Hành vi"],
    [
        ["Bản có sẵn trên thiết bị", "Developer vừa cài build thủ công hoặc cần giữ dữ liệu app", "Không cài đè app trước run"],
        ["Bản build đã tải lên", "Cần xác nhận đúng APK hoặc IPA đã quản lý trong TestPilot", "Cài lại build đã chọn trước khi chạy"],
    ],
    widths=[1.8, 2.55, 2.37],
)

add_section_break(doc)
doc.add_heading("Chạy song song web Android và iOS", level=1)
add_para(
    doc,
    "Trong App Studio, người dùng có thể chọn một nền tảng, hai nền tảng hoặc cả web, Android và iOS. "
    "Hệ thống kiểm tra toàn bộ môi trường trước khi mở driver. Nếu một target bắt buộc chưa sẵn sàng, workflow dừng trước khi chạy "
    "để tránh tình trạng web đã chạy xong còn mobile chưa bắt đầu.",
)
add_flow(doc, ["Preflight tất cả", "Khởi tạo các runner", "Chạy đồng thời", "Gộp dữ liệu", "Tổng hợp verdict"], PALE_GREEN)
add_para(
    doc,
    "Mỗi nền tảng chạy trong một tiến trình riêng: Playwright cho web và Appium cho Android hoặc iOS. Log được gắn nhãn theo nền tảng. "
    "Report, screenshot và video nằm trong run directory riêng. Registry, runtime locator, flake và healing không được ghi đồng thời; "
    "mỗi nhánh lưu phần học được của mình, sau khi tất cả dừng coordinator mới gộp một lần. Cách này giữ tốc độ song song mà tránh lỗi ghi đè dữ liệu.",
)
doc.add_heading("Điều kiện để song song ổn định", level=2)
add_bullets(doc, [
    "Android xuất hiện trong adb devices và ở trạng thái device, không phải unauthorized hoặc offline.",
    "iPhone đã trust máy Mac; devicectl hoặc simulator nhận diện được thiết bị; tunnel và WebDriverAgent sẵn sàng khi cấu hình yêu cầu.",
    "Appium server chấp nhận nhiều session; nhiều thiết bị cùng hệ điều hành phải có systemPort hoặc wdaLocalPort riêng.",
    "Build native phù hợp environment và đã bật inspectability nếu ứng dụng dùng WebView.",
    "Máy local có đủ CPU, RAM và băng thông cho trình duyệt cùng nhiều session Appium.",
])
doc.add_heading("Kiến trúc Local Runner", level=2)
add_flow(doc, ["Feature Gherkin", "Intent Binder", "Resolver", "Playwright hoặc Appium", "Run Report"])

add_section_break(doc)
doc.add_heading("Điểm mạnh và hạn chế của Local Runner", level=1)
add_table(
    doc,
    ["Điểm mạnh", "Giá trị thực tế"],
    [
        ["Chi phí hạ tầng thấp", "Tận dụng máy và thiết bị sẵn có, không tính phút cloud"],
        ["Phản hồi nhanh", "Không chờ queue, upload và cấp phát thiết bị"],
        ["Dễ debug", "Có thể nhìn thiết bị, mở headed browser, xem log trực tiếp và chạy lại tag nhỏ"],
        ["Kiểm soát build và dữ liệu", "Chọn giữ app đang cài hoặc cài lại build upload cho từng run"],
        ["Chạy đa nền tảng", "Một feature chạy riêng lẻ hoặc đồng thời trên web Android và iOS"],
    ],
    widths=[2.05, 4.67],
)
doc.add_heading("Hạn chế", level=2)
add_bullets(doc, [
    ("Phải quản lý thiết bị  ", "Cáp, pin, USB debugging, trust, simulator, driver và OS update đều thuộc trách nhiệm của team."),
    ("Độ phủ phụ thuộc thiết bị đang có  ", "Một máy Android và một iPhone không đại diện cho toàn bộ thị trường."),
    ("Tài nguyên host là giới hạn chung  ", "Nhiều session có thể cạnh tranh CPU, RAM, cổng Appium và băng thông."),
    ("Trạng thái thiết bị dễ làm test nhiễu  ", "Popup quyền, session đăng nhập, dữ liệu cũ và app version khác nhau có thể làm kết quả lệch."),
    ("Không thay thế kiểm thử trên cloud  ", "Local phù hợp debug và xác nhận nhanh, còn độ phủ model và OS vẫn cần Device Farm."),
])
doc.add_heading("AI hỗ trợ Local Runner", level=2)
add_bullets(doc, [
    "Ưu tiên locator đã biết và matching deterministic trước để giữ tốc độ và khả năng giải thích.",
    "Khi locator thất bại, AI semantic hoặc vision xếp hạng candidate dựa trên label, role, màn hình và quan sát runtime.",
    "AI không được tự bấm một hành động có rủi ro khi candidate mơ hồ; cổng ambiguity và action safety có thể dừng để yêu cầu review.",
    "Healing lưu bằng chứng về locator thay thế, tần suất và kết quả retry; report cho phép duyệt trước khi cập nhật bền vững.",
    "Failure diagnosis tách lỗi sản phẩm, locator, automation, dữ liệu và môi trường để người dùng không sửa nhầm lớp.",
])

# Studio
add_section_break(doc)
add_overline(doc, "Phần ba")
doc.add_heading("App Studio Workflow", level=1)
add_speaker_script(
    doc,
    "Hai phần trước bắt đầu từ testcase đã có. App Studio bắt đầu sớm hơn, từ tài liệu nghiệp vụ và ảnh màn hình. "
    "Nó tạo bản nháp kịch bản, đối chiếu coverage, yêu cầu con người review đúng một điểm, rồi tự chuẩn bị môi trường, chạy, healing và sinh report.",
)
doc.add_heading("Điểm khác biệt", level=2)
add_table(
    doc,
    ["Khía cạnh", "Runner", "App Studio Workflow"],
    [
        ["Điểm bắt đầu", "Feature đã tồn tại", "Tài liệu yêu cầu, ticket, file hoặc nguồn qua MCP"],
        ["Vai trò chính", "Thực thi và thu bằng chứng", "Thiết kế testcase, review, thực thi và phản hồi thành một workflow"],
        ["Con người tham gia", "Chọn run và debug khi lỗi", "Review hoặc chỉnh sửa testcase tại cổng duyệt"],
        ["AI", "Fallback discovery, healing và diagnosis", "Phân tích nguồn, vision, sinh Gherkin, coverage và hỗ trợ runtime"],
        ["Kết quả", "Một hoặc nhiều run report", "Feature được quản lý cùng lịch sử quyết định và các report liên quan"],
    ],
    widths=[1.45, 2.25, 3.02],
    font_size=8.8,
)
doc.add_heading("Flow sử dụng", level=2)
add_flow(doc, ["Chọn nguồn", "AI tạo draft", "Review", "Chạy đa nền tảng", "Report và learning"], PALE_GOLD)

add_section_break(doc)
doc.add_heading("Mười một giai đoạn của App Studio", level=1)
add_table(
    doc,
    ["Giai đoạn", "Hệ thống thực hiện", "Điểm kiểm soát"],
    [
        ["1 Đọc và xác thực tài liệu", "Đọc file hoặc gọi MCP, chuẩn hoá nội dung và media", "Báo rõ nguồn thiếu hoặc không đọc được"],
        ["2 AI phân tích yêu cầu màn hình và element", "Rút yêu cầu, màn hình, element và dùng vision khi có ảnh", "Không tự bịa chi tiết ngoài nguồn"],
        ["3 Cập nhật element registry", "Tái sử dụng element và locator đã có, thêm intent mới", "Merge thay vì ghi đè dữ liệu người dùng"],
        ["4 Sinh bộ testcase", "Sinh Gherkin theo vocabulary đóng", "Loại trùng và giữ traceability"],
        ["5 Chuẩn hoá và bind step", "Biến câu tự nhiên thành intent có thể chạy", "Step không bind được phải sửa hoặc review"],
        ["6 Chờ duyệt hoặc chỉnh sửa", "Hiển thị scenario, coverage và action đề xuất", "Con người duyệt, từ chối, sửa hoặc trả lời câu hỏi"],
        ["7 Chuẩn bị môi trường automation", "Preflight web Android iOS và build", "Không mở driver nếu target bắt buộc chưa sẵn sàng"],
        ["8 Chạy các kịch bản đã duyệt", "Chạy riêng lẻ hoặc song song theo nền tảng đã chọn", "Log trực tiếp và nút dừng tác động mọi nhánh"],
        ["9 Healing và chạy lại locator", "Retry chỉ với lỗi locator theo policy", "Không retry assertion nghiệp vụ để tránh che bug"],
        ["10 Sinh report ảnh và video", "Gắn media vào đúng testcase và mốc thời gian", "Phân biệt pass known issue fail và chưa chạy"],
        ["11 Hoàn tất workflow", "Lưu trạng thái, runDirs và lịch sử", "Có thể chạy lại workflow fail mà giữ testcase đã duyệt"],
    ],
    widths=[1.55, 3.0, 2.17],
    font_size=8.2,
)

add_section_break(doc)
doc.add_heading("AI trong App Studio", level=1)
add_para(
    doc,
    "AI được dùng theo từng lớp, không thay thế toàn bộ hệ thống bằng một lời gọi model. Các bước xác định như parse Gherkin, bind intent, "
    "preflight, chạy driver, upload AWS và tính verdict vẫn do code thực hiện. AI tập trung vào những phần cần hiểu ngôn ngữ, hình ảnh "
    "hoặc suy luận candidate.",
)
add_table(
    doc,
    ["Lớp AI", "Dữ liệu đầu vào", "Kết quả hỗ trợ", "Giới hạn"],
    [
        ["Phân tích tài liệu", "Yêu cầu, ticket, nội dung MCP", "Requirement, màn hình, element và câu hỏi còn mơ hồ", "Không tự quyết định nghiệp vụ thiếu dữ kiện"],
        ["Vision", "Ảnh màn hình đi kèm nguồn", "Bổ sung hiểu biết về control, label và bố cục", "Không thay thế locator runtime đã verify"],
        ["Sinh testcase", "Requirement và registry hiện có", "Gherkin ngắn gọn, coverage và traceability", "Bị giới hạn bởi vocabulary có thể thực thi"],
        ["Chuẩn hoá action", "Step tự nhiên chưa bind được", "Action tái sử dụng hoặc chuỗi intent đề xuất", "Cần review nếu phrase template không khớp"],
        ["Element discovery", "UI quan sát tại runtime", "Candidate locator khi deterministic search thất bại", "Phải qua confidence, ambiguity, safety và verification"],
        ["Diagnosis", "Failure, step, screenshot và log", "Giải thích nguyên nhân và hành động tiếp theo", "Rule engine giữ quyền phân loại chính"],
    ],
    widths=[1.28, 1.55, 2.25, 1.64],
    font_size=8.15,
)
doc.add_heading("Điểm mạnh", level=2)
add_bullets(doc, [
    "Tự động hoá liền mạch từ nguồn yêu cầu đến feature, execution và report.",
    "Con người tập trung vào quyết định nghiệp vụ tại cổng review thay vì lặp lại thao tác kỹ thuật.",
    "Tái sử dụng element registry, action và lịch sử healing nên hệ thống cải thiện theo các run đã được xác minh.",
    "Coverage được đối chiếu với yêu cầu quan trọng và nêu rõ phần còn thiếu, thay vì chỉ đếm số testcase.",
    "Một workflow giữ lại stage, log, câu trả lời review và report để truy vết vì sao kết quả được tạo ra.",
])
doc.add_heading("Hạn chế", level=2)
add_bullets(doc, [
    "Chất lượng đầu vào quyết định chất lượng testcase; tài liệu mơ hồ vẫn cần con người trả lời.",
    "Model có thể timeout hoặc đề xuất action chưa tái sử dụng chính xác; bước bổ trợ không nên chặn automation khi có thể tiếp tục an toàn.",
    "AI không tạo được locator đáng tin cậy nếu UI không quan sát được, đặc biệt khi WebView production không bật inspectability.",
    "Không nên chạy thẳng Device Farm khi testcase còn đỏ ở local vì vừa chậm vừa phát sinh chi phí.",
    "Test data và auth fixture cần được quản lý riêng; workflow không thể biến dữ liệu không ổn định thành một test ổn định.",
])

add_section_break(doc)
doc.add_heading("Kịch bản demo đề xuất", level=1)
add_table(
    doc,
    ["Thời gian", "Thao tác", "Điểm cần nói"],
    [
        ["0 đến 2 phút", "Mở dashboard và giới thiệu ba khu vực", "Ba luồng dùng chung feature registry executor và report"],
        ["2 đến 9 phút", "Mở Device Farm, chọn project pool và một run đã có", "AWS nhận app package testspec; mỗi device là một job; artifact được kéo về"],
        ["9 đến 16 phút", "Mở Kịch bản rồi Local Runner", "Chọn tag environment app source; xem preflight log và report"],
        ["16 đến 20 phút", "Chọn web Android iOS trong Studio", "Các nhánh chạy đồng thời; dữ liệu dùng chung chỉ gộp một lần"],
        ["20 đến 28 phút", "Tạo workflow từ tài liệu và dừng ở review", "AI phân tích, sinh testcase và coverage; con người duyệt trước khi chạy"],
        ["28 đến 32 phút", "Mở report hoàn chỉnh", "Phân biệt pass known issue fail flaky; ảnh và video nằm trong từng testcase"],
        ["32 đến 35 phút", "Tổng kết và hỏi đáp", "Chọn runner theo mục tiêu, không theo thói quen"],
    ],
    widths=[1.05, 2.45, 3.22],
    font_size=8.7,
)
doc.add_heading("Dữ liệu demo nên chuẩn bị", level=2)
add_bullets(doc, [
    "Một feature ngắn có case pass, known issue và fail có ảnh hoặc video rõ ràng.",
    "Một APK hoặc IPA đúng environment; nếu demo hybrid phải là build bật inspectability.",
    "Một device pool nhỏ gồm hai đến ba thiết bị có khác biệt model hoặc OS dễ giải thích.",
    "Một workflow đã chạy xong để mở report ngay nếu model hoặc AWS mất thời gian.",
    "Một tài liệu nghiệp vụ có ảnh màn hình và ít nhất một điểm cần review để minh hoạ vai trò con người.",
])
doc.add_heading("Phương án dự phòng khi demo trực tiếp", level=2)
add_bullets(doc, [
    "AWS SSO hết hạn: đăng nhập lại và mở run cũ; giải thích farm pull có thể nối lại việc thu kết quả.",
    "Thiết bị local mất kết nối: dùng preflight để chứng minh hệ thống chặn trước khi chạy, sau đó mở report đã có.",
    "Model timeout: cho thấy workflow ghi cảnh báo và tiếp tục nếu đó là bước bổ trợ, không giả vờ rằng AI luôn sẵn sàng.",
])

add_section_break(doc)
doc.add_heading("Bảng lựa chọn nhanh", level=1)
add_table(
    doc,
    ["Tiêu chí", "Device Farm Runner", "Local Runner", "App Studio Workflow"],
    [
        ["Mục tiêu", "Độ phủ thiết bị", "Debug và phản hồi nhanh", "Tự động hoá toàn vòng đời"],
        ["Điểm bắt đầu", "Feature và app build", "Feature đã duyệt", "Tài liệu nghiệp vụ"],
        ["Web", "Không chạy Playwright", "Có", "Có qua Local Runner"],
        ["Android và iOS", "Thiết bị AWS", "Máy thật hoặc simulator local", "Điều phối local và có thể handoff farm"],
        ["Tốc độ phản hồi", "Chậm hơn do queue và upload", "Nhanh nhất", "Phụ thuộc bước AI review và target chạy"],
        ["Chi phí trực tiếp", "Theo phút thiết bị", "Thấp", "Theo model và runner được chọn"],
        ["AI nổi bật", "Discovery healing diagnosis", "Discovery healing diagnosis", "Sinh testcase coverage vision và toàn bộ hỗ trợ runtime"],
        ["Con người", "Chọn pool và đọc bằng chứng", "Chọn scope và debug", "Review testcase và quyết định điểm mơ hồ"],
    ],
    widths=[1.15, 1.82, 1.72, 2.03],
    font_size=8.05,
)
doc.add_heading("Khuyến nghị sử dụng", level=2)
add_numbered(doc, [
    "Dùng App Studio để tạo hoặc cập nhật testcase từ nguồn yêu cầu và hoàn tất review.",
    "Chạy Local Runner trên scope nhỏ cho đến khi lỗi automation, môi trường và dữ liệu đã được xử lý.",
    "Chạy song song web Android iOS khi cần phản hồi chéo nền tảng và các target đều sẵn sàng.",
    "Đưa suite ổn định lên Device Farm để mở rộng model và OS trước release.",
    "Đọc report theo testcase và evidence; không dùng duy nhất trạng thái xanh hoặc đỏ để kết luận nguyên nhân.",
])
doc.add_heading("Câu kết đề xuất", level=2)
add_speaker_script(
    doc,
    "TestPilot giữ một ngôn ngữ kiểm thử chung nhưng cho phép chọn mức tự động hoá phù hợp với từng thời điểm. "
    "Local Runner tối ưu vòng lặp sửa lỗi, Device Farm mở rộng độ phủ, còn App Studio biến tài liệu và quyết định review thành một quy trình có thể chạy và truy vết.",
)

add_section_break(doc)
doc.add_heading("Câu hỏi thường gặp", level=1)
faq = [
    ("Có thể dùng một testcase cho cả web Android và iOS không", "Có, nếu testcase dùng intent chung và element registry có locator phù hợp cho từng nền tảng. Scenario cũng có thể gắn phạm vi nền tảng khi hành vi thật sự khác nhau."),
    ("AI có tự sửa locator và làm test xanh bằng mọi giá không", "Không. Healing phải qua verification và policy. Assertion nghiệp vụ không được retry như lỗi locator, còn đề xuất bền vững cần review."),
    ("Known issue có làm workflow fail không", "Known issue vẫn xuất hiện rõ trong report nhưng không được tính như một regression mới làm đỏ suite. Khi sản phẩm đã sửa và case pass, report nhắc người dùng gỡ nhãn."),
    ("Tại sao không chạy mọi thứ trên Device Farm", "Device Farm phù hợp độ phủ nhưng tốn thời gian và chi phí. Web Playwright cũng không chạy trong sản phẩm native Device Farm. Local vẫn là nơi debug hiệu quả nhất."),
    ("Nếu AWS chạy xong nhưng giao diện mất kết nối thì sao", "Run vẫn tồn tại trên AWS. TestPilot có thể nối lại theo ARN và tải report cùng artifact bằng luồng farm pull."),
    ("Vì sao mobile local chạy được nhưng farm lại không thấy WebView", "Build local có thể bật inspectability hoặc đã có chromedriver phù hợp, trong khi build gửi farm thì không. Đây là khác biệt build và môi trường, không phải chỉ là khác thiết bị."),
    ("Song song có làm hỏng registry không", "Không theo thiết kế hiện tại. Mỗi nhánh ghi learnings vào thư mục riêng và coordinator gộp đúng một lần sau khi các nhánh dừng."),
]
for question, answer in faq:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(question)
    set_run_font(r, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Inches(0.18)
    rr = p2.add_run(answer)
    set_run_font(rr)

add_section_break(doc)
doc.add_heading("Checklist trước buổi trình bày", level=1)
add_table(
    doc,
    ["Kiểm tra", "Đạt khi"],
    [
        ["UI production", "Server chạy và mở được cổng 4300"],
        ["AWS", "Credential hợp lệ đủ thời gian region us west 2 và project hiển thị"],
        ["Farm build", "APK hoặc IPA test package và testspec đúng phiên bản"],
        ["Android local", "ADB nhận thiết bị Appium sẵn sàng và build đúng environment"],
        ["iOS local", "Thiết bị đã trust tunnel và WDA sẵn sàng hoặc simulator đã boot"],
        ["Web", "Base URL truy cập được và account test đăng nhập được"],
        ["Hybrid", "WebView inspectable và chromedriver tương thích"],
        ["Kịch bản", "Scenario demo đã duyệt tags rõ và known issue đúng content hash"],
        ["Bằng chứng", "Có sẵn report với screenshot video chapter và log dễ giải thích"],
        ["Dự phòng", "Có một farm run một local run và một workflow hoàn chỉnh để mở lại"],
    ],
    widths=[1.65, 5.07],
    font_size=9,
)
doc.add_heading("Nguồn kỹ thuật trong repository", level=2)
add_bullets(doc, [
    "README.md và ARCHITECTURE.md mô tả luồng tổng thể và các quyết định kiến trúc.",
    "src/farm/devicefarm.ts và farm/testspec.yml mô tả lifecycle AWS Device Farm.",
    "src/cli/run.ts và src/cli/run-parallel.ts mô tả execution local và gộp dữ liệu song song.",
    "src/genspec/pipeline.ts mô tả pipeline từ nguồn tài liệu đến Gherkin và coverage.",
    "src/discovery, src/runtime, src/execution và src/healing chứa discovery, executor, retry và healing.",
    "src/report/html.ts và giao diện History hoặc Farm Detail trình bày report cùng media.",
])

add_footer(doc)

# Prevent accidental blue theme colors on all headings and title.
for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name in {"Title", "Heading 1", "Heading 2", "Heading 3"}:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

doc.core_properties.title = "Giới thiệu TestPilot qua ba luồng sử dụng"
doc.core_properties.subject = "Tài liệu trình bày Device Farm Runner Local Runner và App Studio Workflow"
doc.core_properties.author = "TestPilot"
doc.core_properties.keywords = "TestPilot Device Farm Local Runner App Studio AI automation"
doc.save(OUT)
print(OUT)
